#!/usr/bin/env python3
"""
MSPO resource availability + resource-to-quote-line matching (v2)

Drop three files into one folder (e.g. a SharePoint/OneDrive-synced library) and run the script:
    Resource_Plan*.xlsx         supply  (PRP)  - what each resource is booked on, man-days per month
    Quote_Line*.xlsx            demand  (QL)  - pipeline lines, man-days per month
    Position_Name_Mapping.xlsx  sheet "Standard Role" (target roles) + "Position Name Mapping" (labelled examples)

    python mspo_resource_matching_v2.py --input-dir "C:/Users/me/SharePoint/MSPO"          # run once
    python mspo_resource_matching_v2.py --input-dir "C:/Users/me/SharePoint/MSPO" --watch 300   # re-run when files change

Date rules (all dates are 1st of a month):
    Resource end date / "Available From" = month AFTER the last month that has man-days   (last MDs Aug -> 01 Sep)
    Quote-line "Start" date              = FIRST month that has man-days                  (first MDs Sep -> 01 Sep)
    A resource can fill a line when Available From <= Start (and not idle longer than --max-idle months).

Matching order (availability first, then position, then skill):
    0. availability filter   - only resources free inside the rolling window; only lines starting inside it
    T1 standardised position + same grade
    T2 skill (2nd-level check, used when position names do not match / are unavailable) + same grade
    T3 standardised position + grade +/- tolerance
    T4 skill + grade +/- tolerance
Output order: MATCHED, then UNMATCHED RESOURCE, then UNMATCHED PROJECT.
"""
import argparse
import re
import sys
import time
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.optimize import linear_sum_assignment
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# =========================================================================== CONFIG (edit here)
FY_MONTHS = ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]
MONTH_NUM = {m: i for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], 1)}
FULL_MD = 18.0                      # man-days in a month that equals 1 FTE
GRADE_BAND = {3: "B1", 4: "B1", 5: "B2", 6: "B2", 7: "B3", 8: "B4", 9: "B5", 10: "B5", 11: "B5", 12: "B5"}
BIG = 1e6
UNMAPPED = "UNMAPPED - FOR REVIEW"

# Labels typed in the mapping sheet that are not spelled like the "Standard Role" list -> canonical standard role
LABEL_ALIASES = {"service delivery manager": "SDM", "service delivery director": "SDD",
                 "solution architect": "Technical Architect", "operation manager": "SDM"}
# Only these roles read a trailing number as a level ("Cloud Engineer 1" == "Cloud Engineer L1").
# For other roles a trailing number is an instance count (ITSO 2, DC Ops 7).
LEVEL_DIGIT_BASES = {"Cloud Engineer"}

# Ordered rules: first match wins. Left side must be a base name that exists in the "Standard Role" sheet.
ROLE_RULES = [
    ("IT Security Officer", r"\bitso\b|\bsito\b|it security officer|security officer|\biso engineer|^security\s*-\s*l\d"),
    ("IT Security Mgr", r"\bciso\b|security manager|security mgr"),
    ("SDD", r"\bsdd\b|service delivery director|delivery director"),
    ("Senior Project Manager", r"senior project manager|sr\.? project manager"),
    ("Application PM", r"application pm|application project manager"),
    ("Project Manager", r"project manager|project lead|\bpm\b|project mgr"),
    ("SDM", r"\bsdm\b|service delivery (manager|lead)|\bsdl\b|operations? manager|delivery manager|delivery lead|"
            r"dc manager|incident manager|command cent(re|er) manager|cloud technical manager|cloud shared service manager"),
    ("Service Desk Mgr", r"service ?desk\s*(manager|mgr)"),
    ("Service Desk Lead", r"service ?desk\s*(lead|supervisor)"),
    ("Service Desk Agent", r"service ?desk|help ?desk|response cent(re|er)|\bnoc\b"),
    ("Enterprise Architect", r"enterprise architect"),
    ("Application Architect", r"application architect"),
    ("Cloud Architect", r"cloud arc\w*"),
    ("Application Support Lead", r"app(lication)? (support )?lead"),
    ("Application Support Engineer", r"app(lication)? (support )?engineer|application support|app support"),
    ("Technical Architect", r"arc\w*t\w*ect|\barchitect\b|\bsa\b"),
    ("Security Engineer", r"security engineer|cyber ?ark|\bsoc\b|\biam\b|\bpam\b"),
    ("Cloud DBA", r"cloud dba"),
    ("DBA", r"\bdba\b|database"),
    ("Site Reliability Engineer", r"\bsre\b|site reliability"),
    ("FinOps Engineer", r"finops"),
    ("M365 Engineer", r"m365|microsoft 365|o365|office 365|exchange"),
    ("SCCM Engineer", r"sccm|intune|endpoint manager"),
    ("Cloud Engineer", r"cloud|\bgcc\b|\baws\b|azure|\bgcp\b"),
    ("Mobile/Desktop Engineer", r"desktop|mobile|\bdex\b|end ?user|workplace|tech refresh"),
    ("Server Engineer", r"wintel|windows|linux|unix|server|system(s)? engineer|\bsys eng|middleware|vmware|virtuali[sz]ation"),
    ("Network Engineer", r"network|firewall|\bf5\b|load balanc"),
    ("Backup/Storage Engineer", r"storage|back-?up"),
    ("Tools Engineer (Observability)", r"observab|monitoring|tools engineer|\bapm\b"),
    ("IT Service Mgt Engineer", r"itsm|service management|change engineer|servicenow"),
    ("DCOps", r"dc ?op|data cent(re|er)|it ?/? ?(sr )?admin|asset (admin|officer|manager)|\badmin\b|\bdc\b"),
    ("AI Developer", r"\bai\b|machine learning|genai|\bllm\b"),
    ("Frontend Developer", r"front ?end|ui developer"),
    ("Data Engineer", r"data engineer|\betl\b"),
    ("Data Analyst", r"data analyst|\bmis\b"),
    ("Functional Analyst", r"functional analyst"),
    ("Business Analyst", r"business analyst|\bba\b"),
    ("Tester/QA", r"tester|\bqa\b|\buat\b|testing"),
    ("Full Stack Engineer", r"full ?stack|developer|\bjava\b|programmer"),
]

# Skill keyword families (Skill column = ';' separated competencies / certifications)
SKILL_FAMILIES = {
    "SECURITY": ["cyber", "security", "cissp", "cism", "cisa", "siem", "dlp", "edr", "vulnerab", "ccsp", "audit and compliance"],
    "NETWORK": ["network", "cisco", "ccna", "ccnp", "f5 ", "juniper", "checkpoint", "fortinet", "palo alto", "router", "switch"],
    "DATABASE": ["database", "oracle", "sql", "dba", "mysql", "postgres", "mongodb", "redis", "ocp"],
    "CLOUD": ["cloud", "aws", "azure", "gcp", "kubernetes", "docker", "devops", "terraform", "ci/cd"],
    "INFRA_OS": ["linux", "unix", "windows", "wintel", "active directory", "vmware", "vsphere", "storage", "backup",
                 "solaris", "data center", "netapp", "san "],
    "SERVICE_MGMT": ["itil", "servicenow", "itsm", "service desk", "incident", "service management", "remedy", "bmc"],
    "PROJECT_MGMT": ["project management", "pmp", "scrum", "agile", "pmbok", "primavera"],
    "WORKPLACE": ["desktop", "endpoint", "m365", "microsoft 365", "digital workplace", "outlook", "hardware", "end user"],
    "DATA_ANALYTICS": ["power bi", "analytics", "data analysis", "business intelligence", "rpa", "uipath"],
}
# skill family -> base role used ONLY when the position name gives no mapping (2nd-level check)
FAMILY_TO_BASE = {"SECURITY": "IT Security Officer", "NETWORK": "Network Engineer", "DATABASE": "DBA", "CLOUD": "Cloud Engineer",
                  "INFRA_OS": "Server Engineer", "SERVICE_MGMT": "IT Service Mgt Engineer", "PROJECT_MGMT": "Project Manager",
                  "WORKPLACE": "Mobile/Desktop Engineer", "DATA_ANALYTICS": "Data Analyst"}
BASE_TO_FAMILY = {"IT Security Officer": "SECURITY", "IT Security Mgr": "SECURITY", "Security Engineer": "SECURITY",
                  "Network Engineer": "NETWORK", "DBA": "DATABASE", "Cloud DBA": "DATABASE",
                  "Cloud Engineer": "CLOUD", "Cloud Architect": "CLOUD", "FinOps Engineer": "CLOUD", "Site Reliability Engineer": "CLOUD",
                  "Server Engineer": "INFRA_OS", "Backup/Storage Engineer": "INFRA_OS", "DCOps": "INFRA_OS",
                  "IT Service Mgt Engineer": "SERVICE_MGMT", "Service Desk Mgr": "SERVICE_MGMT", "Service Desk Lead": "SERVICE_MGMT",
                  "Service Desk Agent": "SERVICE_MGMT", "SDM": "SERVICE_MGMT", "SDD": "PROJECT_MGMT",
                  "Project Manager": "PROJECT_MGMT", "Senior Project Manager": "PROJECT_MGMT", "Application PM": "PROJECT_MGMT",
                  "Mobile/Desktop Engineer": "WORKPLACE", "M365 Engineer": "WORKPLACE", "SCCM Engineer": "WORKPLACE",
                  "Data Analyst": "DATA_ANALYTICS", "Data Engineer": "DATA_ANALYTICS", "Business Analyst": "DATA_ANALYTICS"}


# =========================================================================== taxonomy (from Position_Name_Mapping.xlsx)
class Taxonomy:
    def __init__(self, roles, redirects, examples):
        self.roles = roles                                  # canonical standard roles
        self.canon = {r.lower(): r for r in roles}
        self.redirects = redirects                          # "SR/IT Admin" -> "DCOps"
        self.parse, self.levels, self.role_of = {}, {}, {}
        for r in roles:
            m = re.match(r"^(.*?)\s+L(\d)\b\s*(\(.*\))?$", r)
            base, lvl = ((m.group(1) + (" " + m.group(3) if m.group(3) else "")).strip(), f"L{m.group(2)}") if m else (r, "")
            self.parse[r] = (base, lvl)
            self.role_of[(base, lvl)] = r
            if lvl:
                self.levels.setdefault(base, []).append(lvl)
        for b in self.levels:
            self.levels[b] = sorted(self.levels[b])
        self.examples = examples                            # norm_key(position name) -> standard role

    def role_name(self, base, level):
        """Standard-form name for base+level. Uses the Standard Role sheet string when it exists; otherwise builds it in the
        same convention (e.g. Cloud Engineer L3 when the sheet only lists L1/L2)."""
        r = self.role_of.get((base, level))
        if r or not level:
            return r or self.role_of.get((base, ""))
        m = re.match(r"^(.*?)(\s*\(.*\))?$", base)
        return f"{m.group(1)} {level}{m.group(2) or ''}"

    def canonical(self, label):
        k = str(label).strip().lower()
        k = LABEL_ALIASES.get(k, k)
        k = self.redirects.get(k.lower(), k)
        return self.canon.get(str(k).lower())


def norm_key(s):
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def load_taxonomy(path):
    x = pd.ExcelFile(path)
    sr = x.parse("Standard Role").iloc[:, 0].dropna().astype(str).str.strip()
    roles, redirects = [], {}
    for r in sr:
        m = re.match(r"^(.*?)\s*\(use\s+(.+?)\)\s*$", r, re.I)
        if m:
            redirects[m.group(1).strip().lower()] = m.group(2).strip()
        else:
            roles.append(r)
    tax = Taxonomy(roles, redirects, {})
    m = x.parse("Position Name Mapping")
    m = m[m["Standard Role"].notna() & m["Position Name"].notna()]
    for name, label in zip(m["Position Name"], m["Standard Role"]):
        role = tax.canonical(label)
        if role:
            tax.examples[norm_key(name)] = role
    return tax


# =========================================================================== step 1: clean + standardise
def clean_position_text(raw):
    s = "" if pd.isna(raw) else str(raw)
    s = re.sub(r"\[[^\]]*\]|\([^)]*\)", " ", s)
    s = re.sub(r"^\s*\d+\s*[.\-)]+\s*", "", s)
    s = re.sub(r"\bHC?R[\s-]?\d+\b|\bHRC\b", " ", s, flags=re.I)
    s = re.sub(r"#\s*\d+", " ", s)
    return re.sub(r"\s+", " ", s).strip(" -–")


def level_text(raw):
    s = "" if pd.isna(raw) else str(raw)
    s = re.sub(r"\bHC?R[\s-]?\d+\b|\bHRC\b|#\s*\d+", " ", s, flags=re.I)
    s = re.sub(r"^\s*\d+\s*[.\-)]+\s*", "", s)
    s = re.sub(r"[\[\](){}]", " ", s)
    return re.sub(r"\s+", " ", s).strip(" -–")


def rule_base(clean, tax):
    for base, rx in ROLE_RULES:
        if base in tax.parse or base in tax.levels or (base, "") in tax.role_of:
            if re.search(rx, clean, re.I):
                return base
    return None


def explicit_level(clean, base, tax):
    """(level, source) read from the name. Only for roles that have levels in the Standard Role sheet."""
    if base not in tax.levels:
        return "", ""
    t = clean.lower()
    m = re.search(r"\bl(?:evel|vl|v)?\s*[-.]?\s*([1-4])(?!\d)", t)
    if m:
        return f"L{m.group(1)}", "name"
    m = re.search(r"(?<![a-z0-9])(iii|ii|i)\s*$", t)
    if m and len(t) > 3:
        return {"i": "L1", "ii": "L2", "iii": "L3"}[m.group(1)], "name"
    if re.search(r"\bsenior\b|\bsr\.?\b|\bsnr\b", t):
        return "L3", "name_title"
    if re.search(r"\bjunior\b|\bjr\b|\bintern\b|\btrainee\b", t):
        return "L1", "name_title"
    if base in LEVEL_DIGIT_BASES:
        m = re.search(r"(?:engineer|cloud)\s*-?\s*([1-3])\s*$", t)
        if m:
            return f"L{m.group(1)}", "name_digit"
    return "", ""


def skill_family(skill):
    if pd.isna(skill):
        return ""
    toks = [t.strip().lower() for t in str(skill).split(";") if t.strip()]
    scores = {f: sum(any(k in t for k in kws) for t in toks) for f, kws in SKILL_FAMILIES.items()}
    fam, top = max(scores.items(), key=lambda kv: kv[1])
    return fam if top > 0 else ""


def resolve_level(base, explicit, grade, tax, lmap):
    lv = tax.levels.get(base, [])
    if not lv:
        return "", ""
    if explicit:
        if explicit in lv:
            return explicit, None                       # keep caller's source label
        return explicit, "name_level_not_in_sheet"
    if len(lv) == 1:
        return lv[0], "single_level"
    g = grade if pd.notna(grade) else 0
    cand = lmap[(lmap["base"] == base) & lmap["level"].isin(lv)] if len(lmap) else lmap
    if len(cand) >= 2 and cand["rows"].sum() >= 3:
        return cand.iloc[(cand["median_grade"] - g).abs().argsort().iloc[0]]["level"], "grade_inferred"
    idx = 0 if g <= 5 else (1 if g <= 7 else 2)
    return lv[min(idx, len(lv) - 1)], "grade_inferred"


def standardise(df, tax):
    """Adds Std columns; every original column is left untouched."""
    out = df.copy()
    n = len(out)
    clean = out["Position Name"].map(clean_position_text)
    skill = out["Skill"] if "Skill" in out else pd.Series([np.nan] * n, index=out.index)
    base, lvl, lsrc, psrc, role = [None] * n, [""] * n, [""] * n, [""] * n, [None] * n

    # pass 1: exact lookup in the mapping sheet, then rules
    for i, (raw, c) in enumerate(zip(out["Position Name"], clean)):
        hit = tax.examples.get(norm_key(raw))
        if hit:
            role[i], psrc[i] = hit, "mapping_sheet"
            base[i], lvl[i] = tax.parse[hit]
            lsrc[i] = "mapping_sheet" if lvl[i] else ""
            continue
        b = rule_base(c, tax)
        if b:
            base[i], psrc[i] = b, "rule"
            lvl[i], lsrc[i] = explicit_level(level_text(raw), b, tax)

    # grade <-> level evidence from explicit levels (data-driven level inference for names without a level)
    ev = pd.DataFrame({"base": base, "level": lvl, "src": lsrc, "grade": out["Grade"].values})
    ev = ev[ev["src"].isin(["name", "name_digit", "mapping_sheet"]) & (ev["level"] != "")]
    lmap = (ev.groupby(["base", "level"])["grade"].agg(median_grade="median", min_grade="min", max_grade="max", rows="count")
            .reset_index()) if len(ev) else pd.DataFrame(columns=["base", "level", "median_grade", "min_grade", "max_grade", "rows"])

    # pass 3: level resolution for rule hits; unmapped names stay unmapped (skill only gives a suggestion)
    fam = skill.map(skill_family)
    sug = [""] * n
    for i in range(n):
        if base[i] is None:
            psrc[i] = "unmapped"
            b = FAMILY_TO_BASE.get(fam.iloc[i])
            if b and (b in tax.levels or (b, "") in tax.role_of):
                l, _ = resolve_level(b, "", out["Grade"].iloc[i], tax, lmap)
                sug[i] = tax.role_name(b, l) or ""
            continue
        if role[i] is None:
            l, s_ = resolve_level(base[i], lvl[i], out["Grade"].iloc[i], tax, lmap)
            lsrc[i] = lsrc[i] if (s_ is None or lsrc[i] == "name_title") else s_
            lvl[i] = l
            role[i] = tax.role_name(base[i], l)
            if role[i] is None:
                psrc[i] = "unmapped"

    def status(ps, ls):
        if ps == "unmapped":
            return "FOR REVIEW"
        if ls == "name_level_not_in_sheet":
            return "MAPPED - level kept from name (not in Standard Role sheet)"
        if ls == "name_title":
            return "MAPPED - level inferred from title (Senior/Junior)"
        if ls == "grade_inferred":
            return "MAPPED - level inferred from grade"
        return "MAPPED"

    out["Position Name Clean"] = clean
    out["Standardized Position"] = [r if r and ps != "unmapped" else UNMAPPED for r, ps in zip(role, psrc)]
    out["Std Position Source"] = psrc
    out["Std Position Status"] = [status(p_, l_) for p_, l_ in zip(psrc, lsrc)]
    out["In Standard Role Sheet"] = ["Y" if r in tax.canon.values() else ("N" if r != UNMAPPED else "") for r in out["Standardized Position"]]
    out["Skill-Inferred Role (suggestion)"] = sug
    out["Std Role Family"] = [b if b and ps != "unmapped" else "" for b, ps in zip(base, psrc)]
    out["Std Level"] = lvl
    out["Std Level Source"] = lsrc
    out["Grade Band"] = out["Grade"].map(GRADE_BAND).fillna("B?")
    own = fam
    default = out["Std Role Family"].map(BASE_TO_FAMILY).fillna("")
    out["Std Skill Family"] = np.where(own != "", own, default)
    out["Std Skill Family Source"] = np.where(own != "", "skill", np.where(default != "", "role_default", ""))
    return out, lmap


def validate_rules(tax):
    """How well do the rules reproduce the labelled examples in the mapping sheet? (without the exact lookup)"""
    rows = []
    for key, label in tax.examples.items():
        c = clean_position_text(key)
        b = rule_base(c, tax)
        c = level_text(key)
        lb, ll = tax.parse[label]
        lv, _ = explicit_level(c, b, tax) if b else ("", "")
        ok_base = b == lb
        ok_level = (not lv) or (not ll) or lv == ll or (lv not in tax.levels.get(b, []))
        rows.append({"Position Name (example)": key, "Labelled Role": label, "Rule Base": b or "", "Rule Level (explicit)": lv,
                     "Base OK": ok_base, "Level OK": ok_level})
    return pd.DataFrame(rows)


# =========================================================================== step 2: dates / availability
LINE_KEY = ["Project ID", "Quote Line Number", "Position Name", "Grade", "Resource"]


def fy_label(ts):
    fy = ts.year if ts.month >= 4 else ts.year - 1
    return f"FY{fy % 100:02d}"


def add_line_id(df):
    d = df.copy()
    d["Slot"] = d.groupby(LINE_KEY + ["Fiscal Year"], dropna=False).cumcount()
    d["Line ID"] = d[LINE_KEY].fillna("").astype(str).agg("|".join, axis=1) + "|" + d["Slot"].astype(str)
    return d


def month_columns(df):
    return [m for m in FY_MONTHS if m in df.columns]


def to_long(d):
    """One row per (line, month) with real 1st-of-month dates. FY26: Apr-Dec 2026, Jan-Mar 2027."""
    d = d.copy()
    d["FY Start Year"] = 2000 + d["Fiscal Year"].str.extract(r"(\d+)")[0].astype(int)
    long = d.melt(id_vars=["Line ID", "FY Start Year"], value_vars=month_columns(d), var_name="Mon", value_name="MDs")
    yr = long["FY Start Year"] + np.where(long["Mon"].isin(["Jan", "Feb", "Mar"]), 1, 0)
    long["Month"] = pd.to_datetime(dict(year=yr, month=long["Mon"].map(MONTH_NUM), day=1))
    long["MDs"] = long["MDs"].fillna(0.0)
    return long.groupby(["Line ID", "Month"], as_index=False)["MDs"].sum()


def line_table(df):
    """One row per line: Start (first month with MDs), Last Active Month, End Date = Available From (month after last MDs)."""
    d = add_line_id(df)
    long = to_long(d)
    act = long[long["MDs"] > 0].sort_values(["Line ID", "Month"])
    g = act.groupby("Line ID")
    s = pd.DataFrame({"Start Date": g["Month"].min(), "Last Active Month": g["Month"].max(),
                      "Avg MDs": g["MDs"].mean(), "Start MDs": g["MDs"].first(), "Last MDs": g["MDs"].last()})
    s["Covered To"] = long.groupby("Line ID")["Month"].max()
    s["Available From"] = s["Last Active Month"] + pd.DateOffset(months=1)
    s["Availability Basis"] = np.where(s["Last Active Month"] < s["Covered To"], "CONFIRMED_END", "FY_END_PLAN")
    s["Start FY"] = s["Start Date"].map(fy_label)
    s["Available From FY"] = s["Available From"].map(fy_label)
    meta = d.drop_duplicates("Line ID").set_index("Line ID")
    return d, meta.join(s, how="inner")


# =========================================================================== step 3: matching
def month_diff(a, b):
    return (b.year - a.year) * 12 + (b.month - a.month)


def skill_similarity(sup, dem):
    corpus = pd.concat([sup["Skill"].fillna(""), dem["Skill"].fillna("")]).str.lower().str.replace(r"[^a-z0-9+#/ ]", " ", regex=True)
    if corpus.str.strip().eq("").all():
        return np.zeros((len(sup), len(dem))), np.zeros(len(sup), bool), np.zeros(len(dem), bool)
    vec = TfidfVectorizer(ngram_range=(1, 2), stop_words="english").fit(corpus)
    sv, dv = vec.transform(corpus.iloc[:len(sup)]), vec.transform(corpus.iloc[len(sup):])
    return cosine_similarity(sv, dv), sup["Skill"].notna().to_numpy(), dem["Skill"].notna().to_numpy()


def run_match(sup, dem, grade_tol, max_idle, skill_thr):
    sup, dem = sup.reset_index(), dem.reset_index()
    ns, nd = len(sup), len(dem)
    sim, s_has, d_has = skill_similarity(sup, dem)
    gap = np.array([[month_diff(a, b) for b in dem["Start Date"]] for a in sup["Available From"]]) if ns and nd else np.zeros((ns, nd), int)
    gd = np.abs(sup["Grade"].to_numpy(float)[:, None] - dem["Grade"].to_numpy(float)[None, :])
    same = sup["Line ID"].to_numpy()[:, None] == dem["Line ID"].to_numpy()[None, :]
    sp, dp = sup["Standardized Position"].to_numpy(), dem["Standardized Position"].to_numpy()
    s_pos_ok = (sup["Standardized Position"] != UNMAPPED).to_numpy()
    d_pos_ok = (dem["Standardized Position"] != UNMAPPED).to_numpy()
    pos_eq = (sp[:, None] == dp[None, :]) & s_pos_ok[:, None] & d_pos_ok[None, :]
    sf, df_ = sup["Std Skill Family"].to_numpy(), dem["Std Skill Family"].to_numpy()
    fam_eq = (sf[:, None] == df_[None, :]) & (sf[:, None] != "")
    both = s_has[:, None] & d_has[None, :]
    skill_score = np.where(both, sim, np.where(fam_eq, 1.0, 0.0))
    skill_ok = fam_eq & (~both | (sim >= skill_thr))
    fy_pen = (sup["Availability Basis"].to_numpy() == "FY_END_PLAN").astype(float)[:, None] * 0.5
    free_ok = (~same) & (gap >= 0) & (gap <= max_idle)              # availability gate: free on/before start

    tiers = [("T1", "Position + Grade", pos_eq, 0), ("T2", "Skill + Grade", skill_ok, 0),
             ("T3", "Position + Grade (+/-)", pos_eq, grade_tol), ("T4", "Skill + Grade (+/-)", skill_ok, grade_tol)]
    sfree, dfree, pairs = np.ones(ns, bool), np.ones(nd, bool), []
    for tid, label, rule, tol in tiers:
        feas = free_ok & rule & (gd <= tol) & sfree[:, None] & dfree[None, :]
        if not feas.any():
            continue
        cost = np.where(feas, gap + 2.0 * gd + (1.0 - skill_score) + fy_pen, BIG)
        r, c = linear_sum_assignment(cost)
        for i, j in zip(r, c):
            if cost[i, j] < BIG:
                pairs.append((i, j, tid, label, float(skill_score[i, j]), int(gap[i, j]), int(gd[i, j])))
                sfree[i], dfree[j] = False, False
    return sup, dem, pairs, sfree, dfree


S_COLS = ["Project ID", "Quote Line Number", "Customer", "Project Name", "Resource", "Position Name", "Standardized Position",
          "Std Position Source", "Std Position Status", "Skill-Inferred Role (suggestion)", "Grade", "Std Skill Family", "Skill", "Last Active Month", "Available From", "Available From FY",
          "Availability Basis", "Availability Status", "Last MDs"]
D_COLS = ["Project ID", "Quote Line Number", "Customer", "Project Name", "Position Name", "Standardized Position",
          "Std Position Source", "Std Position Status", "Skill-Inferred Role (suggestion)", "Grade", "Std Skill Family", "Skill", "Start Date", "Start FY", "Start MDs"]


def review_flag(*positions):
    return "FOR REVIEW - position name unmapped" if any(p == UNMAPPED for p in positions) else ""


def assemble(sup, dem, pairs, sfree, dfree):
    sp = sup.reindex(columns=S_COLS).add_prefix("PRP | ")
    dp = dem.reindex(columns=D_COLS).add_prefix("QL | ")
    for p in (sp, dp):
        for c in p.columns:
            if c.endswith("| Skill"):
                p[c] = p[c].astype("string").str.slice(0, 160)
    rows = []
    for i, j, tid, label, sc, g, gdiff in sorted(pairs, key=lambda p: (p[2], dem.at[p[1], "Start Date"], p[5])):
        reason = (f"{label}; available {sup.at[i, 'Available From']:%Y-%m-%d} <= start {dem.at[j, 'Start Date']:%Y-%m-%d} "
                  f"(idle {g}m); grade {sup.at[i, 'Grade']:.0f} vs {dem.at[j, 'Grade']:.0f}")
        r = {"Match Status": "MATCHED", "Match Tier": tid, "Match Basis": label, "Skill Score": round(sc, 2),
             "Idle Months": g, "Grade Gap": gdiff, "Match Reason": reason,
             "Review Flag": review_flag(sup.at[i, "Standardized Position"], dem.at[j, "Standardized Position"])}
        r.update(sp.iloc[i].to_dict()); r.update(dp.iloc[j].to_dict()); rows.append(r)
    for i in np.where(sfree)[0]:
        r = {"Match Status": "UNMATCHED RESOURCE", "Match Reason": "no line with compatible position/skill, grade and start date in window",
             "Review Flag": review_flag(sup.at[i, "Standardized Position"])}
        r.update(sp.iloc[i].to_dict()); rows.append(r)
    for j in np.where(dfree)[0]:
        r = {"Match Status": "UNMATCHED PROJECT", "Match Reason": "no available resource with compatible position/skill, grade and date",
             "Review Flag": review_flag(dem.at[j, "Standardized Position"])}
        r.update(dp.iloc[j].to_dict()); rows.append(r)
    cols = ["Match Status", "Review Flag", "Match Tier", "Match Basis", "Skill Score", "Idle Months", "Grade Gap", "Match Reason"]
    res = pd.DataFrame(rows)
    return res[cols + [c for c in res.columns if c not in cols]]


# =========================================================================== step 4: rolling availability forecast
def build_forecast(sup, dem, matches, as_of, horizon, renewal_rate):
    """Monthly (as_of .. as_of+horizon-1) by standardised position: availability, pipeline demand, matched, surplus/gap."""
    months = pd.date_range(as_of, periods=horizon, freq="MS")
    m = matches[matches["Match Status"] == "MATCHED"]
    roles = sorted(set(sup["Standardized Position"]) | set(dem["Standardized Position"]))
    rows = []

    def block(role, s, d, mm):
        cum_av = cum_conf = cum_fy = cum_match = cum_dem = 0.0
        for k, mo in enumerate(months):
            first = k == 0
            in_m = (s["Available From"] <= mo) if first else (s["Available From"] == mo)       # opening stock lands in month 1
            new_c = s[in_m & (s["Availability Basis"] == "CONFIRMED_END")]
            new_f = s[in_m & (s["Availability Basis"] == "FY_END_PLAN")]
            dem_m = d[d["Start Date"] == mo]
            mat_m = mm[mm["QL | Start Date"] == mo]
            cum_conf += len(new_c); cum_fy += len(new_f); cum_dem += len(dem_m); cum_match += len(mat_m)
            cum_av = cum_conf + cum_fy
            rows.append({
                "Month": mo, "Horizon": "0-6 months" if k < 6 else "7-12 months", "Standardized Position": role,
                "New Available - confirmed end": len(new_c), "New Available - FY-end (plan)": len(new_f),
                "New Available FTE": round((new_c["Last MDs"].sum() + new_f["Last MDs"].sum()) / FULL_MD, 2),
                "Cumulative Available (plan)": int(cum_av),
                "Cumulative Available (risk-adjusted)": round(cum_conf + cum_fy * (1 - renewal_rate), 1),
                "Pipeline Demand Starting": len(dem_m), "Pipeline Demand FTE": round(dem_m["Start MDs"].sum() / FULL_MD, 2),
                "Matched Starting": len(mat_m), "Cumulative Matched": int(cum_match),
                "Surplus Resources After Matching": int(cum_av - cum_match),
                "Unfilled Pipeline (cumulative)": int(cum_dem - cum_match),
                "Coverage % (cum. matched / cum. demand)": round(100 * cum_match / cum_dem, 0) if cum_dem else np.nan})

    for role in roles:
        block(role, sup[sup["Standardized Position"] == role], dem[dem["Standardized Position"] == role],
              m[m["PRP | Standardized Position"] == role])
    block("ALL ROLES", sup, dem, m)
    return pd.DataFrame(rows)


def forecast_summary(fc):
    out = []
    for label, k in (("Next 6 months", 6), ("Next 12 months", 12)):
        mo = sorted(fc["Month"].unique())[:k]
        last = fc[fc["Month"] == mo[-1]].set_index("Standardized Position")
        sub = fc[fc["Month"].isin(mo)].groupby("Standardized Position").agg(
            new_conf=("New Available - confirmed end", "sum"), new_fy=("New Available - FY-end (plan)", "sum"),
            demand=("Pipeline Demand Starting", "sum"), matched=("Matched Starting", "sum"))
        d = sub.join(last[["Cumulative Available (plan)", "Cumulative Available (risk-adjusted)", "Surplus Resources After Matching",
                           "Unfilled Pipeline (cumulative)"]])
        d.insert(0, "Window", label)
        out.append(d.reset_index())
    return pd.concat(out, ignore_index=True)


# =========================================================================== IO helpers
def discover(input_dir):
    d = Path(input_dir)
    pick = lambda pat: max(d.glob(pat), key=lambda p: p.stat().st_mtime, default=None)
    found = {"prp": pick("[Rr]esource*[Pp]lan*.xlsx"), "ql": pick("[Qq]uote*[Ll]ine*.xlsx"), "map": pick("[Pp]osition*[Mm]apping*.xlsx")}
    missing = [k for k, v in found.items() if v is None]
    if missing:
        sys.exit(f"Missing file(s) in {d}: {missing}. Expected Resource_Plan*.xlsx, Quote_Line*.xlsx, Position_Name_Mapping*.xlsx")
    return found


def snapshot_as_of(path, override):
    if override:
        ts = pd.Timestamp(override)
    else:
        m = re.search(r"(\d{2})(\d{2})(\d{2})(?!\d)", Path(path).stem)
        try:
            ts = pd.Timestamp(2000 + int(m.group(3)), int(m.group(2)), int(m.group(1)))      # ddmmyy in file name
        except Exception:
            ts = pd.Timestamp.today()
    ts = ts.normalize()
    return ts if ts.day == 1 else (ts + pd.offsets.MonthBegin(1))


def style_xlsx(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = load_workbook(path)
    fills = {"MATCHED": "E2F0D9", "UNMATCHED RESOURCE": "FFF2CC", "UNMATCHED PROJECT": "FCE4D6"}
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200])
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, width + 2), 48)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if hasattr(c.value, "year"):
                    c.number_format = "dd-mmm-yyyy"
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, str) and "FOR REVIEW" in c.value:
                    c.fill, c.font = PatternFill("solid", fgColor="FFC000"), Font(bold=True, color="C00000")
        if ws.title == "Positions For Review" and "Standard Role List" in wb.sheetnames:
            from openpyxl.worksheet.datavalidation import DataValidation
            n_roles = wb["Standard Role List"].max_row
            dv = DataValidation(type="list", formula1=f"='Standard Role List'!$A$2:$A${n_roles}", allow_blank=True)
            ws.add_data_validation(dv)
            col = [c.value for c in ws[1]].index("Standard Role (please fill)") + 1
            dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{max(ws.max_row, 2)}")
        if ws.title == "Matching":
            for row in ws.iter_rows(min_row=2):
                f = fills.get(row[0].value)
                if f:
                    for c in [row[0]] + list(row[2:7]):
                        c.fill = PatternFill("solid", fgColor=f)
    wb.save(path)


# =========================================================================== orchestration
def run(a):
    files = discover(a.input_dir)
    out = Path(a.output_dir or Path(a.input_dir) / "output"); out.mkdir(parents=True, exist_ok=True)
    as_of = snapshot_as_of(files["prp"], a.as_of)
    print(f"files: {', '.join(p.name for p in files.values())}\nplanning window: {as_of:%d-%b-%Y} -> +{a.horizon} months")

    tax = load_taxonomy(files["map"])
    prp_raw, ql_raw = (pd.read_excel(files[k]) for k in ("prp", "ql"))
    prp, lmap_p = standardise(prp_raw, tax)
    ql, lmap_q = standardise(ql_raw, tax)

    # line-level dates + clean csvs (original columns first, new columns after)
    prp_d, prp_lines = line_table(prp)
    ql_d, ql_lines = line_table(ql)
    for clean_df, d in ((prp, prp_d), (ql, ql_d)):
        clean_df["Line ID"] = d["Line ID"].values
    date_cols = ["Start Date", "Last Active Month", "Available From", "Available From FY", "Availability Basis"]
    prp_out = prp.join(prp_lines[date_cols], on="Line ID")
    ql_out = ql.join(ql_lines[date_cols], on="Line ID")

    # availability filter FIRST
    end_of_window = as_of + pd.DateOffset(months=a.horizon)
    s = prp_lines[prp_lines["Available From"] < end_of_window].copy()
    if a.confirmed_only:
        s = s[s["Availability Basis"] == "CONFIRMED_END"]
    s["Availability Status"] = np.select(
        [s["Available From"] <= as_of, s["Available From"] < as_of + pd.DateOffset(months=6)],
        ["Available now", "Rolling off 0-6m"], "Rolling off 7-12m")
    d = ql_lines[(ql_lines["Start Date"] >= as_of) & (ql_lines["Start Date"] < end_of_window)].copy()

    sup, dem, pairs, sfree, dfree = run_match(s, d, a.grade_tol, a.max_idle, a.skill_thr)
    res = assemble(sup, dem, pairs, sfree, dfree)
    fc = build_forecast(sup, dem, res, as_of, a.horizon, a.renewal_rate)
    fsum = forecast_summary(fc)

    # outputs
    val = validate_rules(tax)
    mapping_tbl = (pd.concat([prp, ql])[["Position Name", "Position Name Clean", "Standardized Position", "Std Role Family", "Std Level",
                                         "Std Level Source", "Std Position Source", "Std Position Status", "In Standard Role Sheet"]]
                   .drop_duplicates("Position Name").sort_values(["Standardized Position", "Position Name"]))
    allrows = pd.concat([prp.assign(Source="PRP"), ql.assign(Source="QL")])
    rev = allrows[allrows["Standardized Position"] == UNMAPPED]
    review = (rev.groupby("Position Name").agg(
        **{"PRP rows": ("Source", lambda x: int((x == "PRP").sum())), "QL rows": ("Source", lambda x: int((x == "QL").sum())),
           "Grade(s)": ("Grade", lambda x: ", ".join(str(int(g)) for g in sorted(x.dropna().unique()))),
           "Example Project ID": ("Project ID", "first"), "Example Customer": ("Customer", "first"),
           "Skill family (from Skill)": ("Std Skill Family", lambda x: next((v for v in x if v), "")),
           "Skill-Inferred Role (suggestion)": ("Skill-Inferred Role (suggestion)", lambda x: next((v for v in x if v), ""))})
        .reset_index().sort_values(["PRP rows", "QL rows"], ascending=False))
    review.insert(0, "Review Flag", "FOR REVIEW - position name unmapped")
    review["Standard Role (please fill)"] = ""
    not_in_sheet = sorted(set(allrows.loc[allrows["In Standard Role Sheet"] == "N", "Standardized Position"]))
    bad = [r for r in set(allrows["Standardized Position"]) if r != UNMAPPED and r not in tax.canon.values() and r not in not_in_sheet]
    assert not bad, bad
    summary = pd.DataFrame({"Item": [
        "Planning date (as-of)", "Rolling window (months)", "Resources available in window", "QL lines starting in window",
        "MATCHED", "UNMATCHED RESOURCE", "UNMATCHED PROJECT", "Matches by position+grade (T1)", "Matches by skill+grade (T2)",
        "Matches by position, grade +/- (T3)", "Matches by skill, grade +/- (T4)", "Grade tolerance (+/-)", "Max idle months",
        "FY-end renewal rate (risk-adjusted forecast)", "FOR REVIEW: distinct position names", "FOR REVIEW: PRP rows / QL rows",
        "Rule accuracy vs your labelled examples (role)", "Standardized Position source: mapping_sheet / rule / unmapped (PRP rows)",
        "Names kept with a level not in the Standard Role sheet"],
        "Value": [f"{as_of:%d-%b-%Y}", a.horizon, len(sup), len(dem),
                  int((res['Match Status'] == 'MATCHED').sum()), int((res['Match Status'] == 'UNMATCHED RESOURCE').sum()),
                  int((res['Match Status'] == 'UNMATCHED PROJECT').sum()),
                  *[int((res.get('Match Tier') == t).sum()) for t in ("T1", "T2", "T3", "T4")], a.grade_tol, a.max_idle, a.renewal_rate,
                  len(review), f"{int((prp['Standardized Position'] == UNMAPPED).sum())} / {int((ql['Standardized Position'] == UNMAPPED).sum())}",
                  f"{100 * val['Base OK'].mean():.0f}%",
                  " / ".join(str(int((prp['Std Position Source'] == k).sum())) for k in ("mapping_sheet", "rule", "unmapped")),
                  ", ".join(not_in_sheet) or "none"]})
    csv_kw = dict(index=False, date_format="%Y-%m-%d")
    prp_out.to_csv(out / "resource_plan_clean.csv", **csv_kw)
    ql_out.to_csv(out / "quote_line_clean.csv", **csv_kw)
    mapping_tbl.to_csv(out / "position_mapping_table.csv", **csv_kw)
    val[~(val["Base OK"] & val["Level OK"])].to_csv(out / "mapping_validation_disagreements.csv", **csv_kw)
    review.to_csv(out / "positions_for_review.csv", **csv_kw)
    res.to_csv(out / "matching_result.csv", **csv_kw)
    fc.to_csv(out / "availability_forecast.csv", **csv_kw)
    xlsx = out / "matching_result.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        res.to_excel(xw, sheet_name="Matching", index=False)
        res[res["Match Status"] == "MATCHED"].dropna(axis=1, how="all").to_excel(xw, sheet_name="Matched", index=False)
        res[res["Match Status"] == "UNMATCHED RESOURCE"].dropna(axis=1, how="all").to_excel(xw, sheet_name="Unmatched Resources", index=False)
        res[res["Match Status"] == "UNMATCHED PROJECT"].dropna(axis=1, how="all").to_excel(xw, sheet_name="Unmatched Projects", index=False)
        fsum.to_excel(xw, sheet_name="Forecast 6-12m Summary", index=False)
        fc.to_excel(xw, sheet_name="Availability Forecast", index=False)
        review.to_excel(xw, sheet_name="Positions For Review", index=False)
        mapping_tbl.to_excel(xw, sheet_name="Position Mapping", index=False)
        pd.DataFrame({"Standard Role": tax.roles}).to_excel(xw, sheet_name="Standard Role List", index=False)
    style_xlsx(xlsx)
    print(summary.to_string(index=False))
    print("written to", out.resolve())


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input-dir", default=".", help="folder holding the 3 xlsx files (e.g. synced SharePoint library)")
    ap.add_argument("--output-dir", default=None, help="default: <input-dir>/output")
    ap.add_argument("--as-of", default=None, help="planning date YYYY-MM-DD; default = month after the ddmmyy date in the Resource Plan file name")
    ap.add_argument("--horizon", type=int, default=12, help="rolling window in months (6-12)")
    ap.add_argument("--grade-tol", type=int, default=1, help="grade +/- allowed in tiers T3/T4 (0 = off)")
    ap.add_argument("--max-idle", type=int, default=6, help="max months a free resource may wait for the line start")
    ap.add_argument("--skill-thr", type=float, default=0.15, help="min TF-IDF skill similarity when both sides list skills")
    ap.add_argument("--renewal-rate", type=float, default=0.8, help="assumed share of FY-end lines that get extended (risk-adjusted column only)")
    ap.add_argument("--confirmed-only", action="store_true", help="ignore lines that simply run to fiscal-year end")
    ap.add_argument("--watch", type=int, default=0, help="poll the folder every N seconds and re-run when a file changes")
    a = ap.parse_args()
    if not a.watch:
        return run(a)
    seen = None
    while True:
        files = discover(a.input_dir)
        stamp = tuple(p.stat().st_mtime for p in files.values())
        if stamp != seen:
            seen = stamp
            try:
                run(a)
            except Exception as e:                                     # keep watching after a bad upload
                print("run failed:", e)
        time.sleep(a.watch)


if __name__ == "__main__":
    main()
